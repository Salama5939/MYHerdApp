import docx

doc = docx.Document()

# Document Title
doc.add_heading(
    "Comprehensive Bilingual Administration & Reference Manual: MyHerdApp &"
    " MyHerdFinance",
    level=0,
)
doc.add_heading(
    "الدليل الإداري والمرجعي الشامل ثنائي اللغة: تطبيق MyHerdApp وتطبيق"
    " MyHerdFinance",
    level=1,
)

# Part 1: System Overview & Reference User Profile
doc.add_heading("Part 1: System Overview & Reference User Profile", level=2)
doc.add_heading("الجزء الأول: نظرة عامة على النظام وملف المستخدم المرجعي", level=2)

doc.add_heading("1. Reference User Profile / ملف المستخدم المرجعي", level=3)
doc.add_paragraph(
    "English: Farm Manager / Owner — An agricultural manager or livestock"
    " breeder operating a sheep production enterprise who requires precise"
    " tracking of animal lifecycles, operational logistics, labor performance,"
    " and financial accounting."
)
doc.add_paragraph(
    "Arabic: مدير / مالك المزرعة — مدير زراعي أو مربي مواشي يدير مشروع إنتاج"
    " أغنام يتطلب متابعة دقيقة لدورة حياة الحيوانات، اللوجستيات التشغيلية، أداء"
    " العمالة، والحسابات المالية."
)

doc.add_heading("2. Core Application Aims / أهداف التطبيق الأساسية", level=3)
doc.add_paragraph(
    "MyHerdApp (English): To digitize and streamline the biological and"
    " operational lifecycle of a sheep farm—ranging from herd registration,"
    " birth tracking, and growth monitoring to feed management and custom"
    " butchering/cutting orders."
)
doc.add_paragraph(
    "MyHerdApp (Arabic): رقمية وتبسيط دورة الحياة البيولوجية والتشغيلية لمزرعة"
    " الأغنام — بدءاً من تسجيل القطيع، وتتبع المواليد، ومراقبة النمو، إلى إدارة"
    " الأعلاف وطلبات التقطيع والذبح الخاصة."
)
doc.add_paragraph(
    "MyHerdFinance (English): To track incoming and outgoing financial"
    " transactions, manage cash flow rosters, analyze fattening efficiency, and"
    " maintain a clear ledger of farm economics."
)
doc.add_paragraph(
    "MyHerdFinance (Arabic): تتبع المعاملات المالية الواردة والصادرة، إدارة جدول"
    " التدفق النقدي، تحليل كفاءة التسمين، والحفاظ على دفتر أستاذ واضح لاقتصاديات"
    " المزرعة."
)

# Part 2: MyHerdFinance Detailed Module Manual
doc.add_heading(
    "Part 2: MyHerdFinance Detailed Module-by-Module Operational Manual",
    level=2,
)
doc.add_heading(
    "الجزء الثاني: الدليل التشغيلي المفصل لوحدات تطبيق MyHerdFinance", level=2
)

modules_finance = [
    (
        "1. Home / الرئيسية",
        (
            "Acts as the primary secure landing portal and authentication"
            " gateway requiring authorized credentials to access financial"
            " ledgers."
        ),
        (
            "يعمل بمثابة بوابة الهبوط الأساسية الآمنة وبوابة المصادقة التي تتطلب"
            " بيانات معتمدة للوصول إلى السجلات المالية."
        ),
        (
            "Accessed upon initial startup or when logging into the financial"
            " executive suite."
        ),
        (
            "يتم الوصول إليها عند بدء التشغيل الأولي أو عند تسجيل الدخول إلى"
            " الجناح المالي التنفيذي."
        ),
    ),
    (
        "2. Dashboard / لوحة التحكم",
        (
            "Provides a high-level visual summary of the farm's financial"
            " health, cash standing, and quick financial performance"
            " indicators."
        ),
        (
            "يوفر ملخصاً مرئياً عالي المستوى للصحة المالية للمزرعة، الوضع النقدي،"
            " ومؤشرات الأداء المالي السريعة."
        ),
        (
            "Consulted daily by management for a quick macro-level snapshot of"
            " financial metrics."
        ),
        (
            "يتم استشارته يومياً من قبل الإدارة للحصول على نظرة عامة سريعة على"
            " المقاييس المالية."
        ),
    ),
    (
        "3. Cash Flow Roster / جدول التدفق النقدي",
        (
            "Tracks chronological cash movements, liquidity positions, and"
            " transaction rosters across farm operations."
        ),
        (
            "يتتبع الحركات النقدية الزمنية، مواضع السيولة، وجداول المعاملات عبر"
            " عمليات المزرعة."
        ),
        (
            "Updated and reviewed frequently to ensure sufficient operational"
            " liquidity."
        ),
        ("يتم تحديثه ومراجعته بشكل متكرر لضمان توفر سيولة تشغيلية كافية."),
    ),
    (
        "4. Strategic Capital / رأس المال الاستراتيجي",
        (
            "Manages long-term capital assets, core investments, infrastructure"
            " valuations, and structural financial allocations."
        ),
        (
            "يدير الأصول الرأسمالية طويلة الأجل، الاستثمارات الأساسية، تقييمات"
            " البنية التحتية، والتخصيصات المالية الهيكلية."
        ),
        (
            "Reviewed during strategic quarterly or annual capital investment"
            " evaluations."
        ),
        (
            "يتم مراجعته خلال تقييمات الاستثمار الرأسمالي الاستراتيجية ربع"
            " السنوية أو السنوية."
        ),
    ),
    (
        "5. Operational Outflow / المصاريف التشغيلية",
        (
            "Records and categorizes all farm expenditures, including feed"
            " purchases, veterinary medicines, labor wages, and maintenance"
            " costs."
        ),
        (
            "يسجل ويصنف جميع نفقات المزرعة، بما في ذلك مشتريات الأعلاف، الأدوية"
            " البيطرية، أجور العمالة، وتكاليف الصيانة."
        ),
        ("Updated daily as expenses occur to maintain precise cost control."),
        ("يتم تحديثه يومياً عند حدوث المصروفات للحفاظ على تحكم دقيق في" " التكاليف."),
    ),
    (
        "6. Revenue Inflow / سجل الإيرادات",
        (
            "Logs incoming monetary streams from livestock sales, wool, manure,"
            " meat cutting orders, and other farm commercial activities."
        ),
        (
            "يسجل التدفقات النقدية الواردة من مبيعات المواشي، الصوف، السماد، طلبات"
            " تقطيع اللحوم، والأنشطة التجارية الأخرى للمزرعة."
        ),
        (
            "Updated immediately following any commercial sales transaction or"
            " client payment collection."
        ),
        ("يتم تحديثه فوراً بعد أي معاملة بيع تجارية أو تحصيل مدفوعات من" " العملاء."),
    ),
    (
        "7. Data Bridge / جسر البيانات",
        (
            "Serves as the integration utility connecting operational metrics"
            " and physical inventory movements from MyHerdApp with financial"
            " ledgers in MyHerdFinance."
        ),
        (
            "يعمل بمثابة أداة التكامل التي تربط المقاييس التشغيلية وحركات"
            " المخزون الفيزيائي من تطبيق MyHerdApp بالسجلات المالية في تطبيق"
            " MyHerdFinance."
        ),
        (
            "Utilized periodically (weekly/monthly) during cross-app"
            " synchronization audits and inventory cost-valuation updates."
        ),
        (
            "يُستخدم دورياً (أسبوعياً/شهرياً) أثناء عمليات تدقيق المزامنة بين"
            " التطبيقات وتحديثات تقييم تكلفة المخزون."
        ),
    ),
    (
        "8. Farm Analytics / التحليلات المالية",
        (
            "Performs deep financial evaluations, cost-benefit analyses, and"
            " economic trend forecasting for the farm enterprise."
        ),
        (
            "يجري تقييمات مالية عميقة، تحليلات التكلفة والعائد، والتنبؤ"
            " بالاتجاهات الاقتصادية لمشروع المزرعة."
        ),
        (
            "Consulted during monthly financial closing and strategic business"
            " planning sessions."
        ),
        (
            "يتم استشارته خلال جلسات الإقفال المالي الشهري وتخطيط الأعمال"
            " الاستراتيجي."
        ),
    ),
    (
        "9. Fattening Efficiency / كفاءة التسمين",
        (
            "Calculates the economic return on investment (ROI) for feeding and"
            " growth batches, linking feed input costs to final weight gain"
            " values."
        ),
        (
            "يحسب العائد الاقتصادي على الاستثمار (ROI) لدفعات التغذية والنمو،"
            " رابطاً تكاليف مدخلات الأعلاف بقيم زيادة الوزن النهائية."
        ),
        (
            "Reviewed at the completion of a fattening cycle or batch marketing"
            " window."
        ),
        ("يتم مراجعته عند اكتمال دورة التسمين أو نافذة تسويق الدفعات."),
    ),
    (
        "10. Cash Flow Forecast / توقعات التدفق النقدي",
        (
            "Projects future cash positions, anticipated revenues, and"
            " projected liabilities based on historical operational data."
        ),
        (
            "يتوقع المراكز النقدية المستقبلية، الإيرادات المتوقعة، والالتزامات"
            " المقدرة بناءً على البيانات التشغيلية التاريخية."
        ),
        (
            "Consulted before major purchasing or marketing decisions to plan"
            " upcoming financial needs."
        ),
        (
            "يتم استشارته قبل اتخاذ قرارات الشراء أو التسويق الكبرى لتخطيط"
            " الاحتياجات المالية القادمة."
        ),
    ),
]

for title, func_en, func_ar, time_en, time_ar in modules_finance:
    doc.add_heading(title, level=3)
    doc.add_paragraph(f"Function (English): {func_en}")
    doc.add_paragraph(f"الوظيفة (عربي): {func_ar}")
    doc.add_paragraph(f"Operational Timing (English): {time_en}")
    doc.add_paragraph(f"التوقيت التشغيلي (عربي): {time_ar}")

# Part 3: MyHerdApp Detailed Module Manual
doc.add_heading(
    "Part 3: MyHerdApp Detailed Module-by-Module Operational Manual", level=2
)
doc.add_heading("الجزء الثالث: الدليل التشغيلي المفصل لوحدات تطبيق MyHerdApp", level=2)

modules_app = [
    (
        "1. Strategic Metrics / المقاييس الاستراتيجية",
        (
            "Provides a high-level executive dashboard tracking core"
            " operational key performance indicators (KPIs) and long-term farm"
            " health trends."
        ),
        (
            "يوفر لوحة تحكم تنفيذية عالية المستوى تتبع مؤشرات الأداء الرئيسية"
            " التشغيلية (KPIs) واتجاهات صحة المزرعة طويلة الأجل."
        ),
        (
            "Reviewed by farm managers and owners during weekly or monthly"
            " managerial meetings."
        ),
        (
            "يتم مراجعته من قبل مديري المزارع والمالكين خلال الاجتماعات الإدارية"
            " الأسبوعية أو الشهرية."
        ),
    ),
    (
        "2. Active Herd Registry / سجل القطيع النشط",
        (
            "Serves as the foundational database for every living animal,"
            " capturing unique ear tag/RFID identifiers, birth dates, gender,"
            " genealogy, and operational status."
        ),
        (
            "يعمل بمثابة قاعدة البيانات الأساسية لكل حيوان حي، مسجلاً أرقام الأذن"
            " التعريفية الفريدة، تواريخ الولادة، الجنس، النسب، والحالة"
            " التشغيلية."
        ),
        (
            "Used on-the-ground immediately when new livestock are tagged,"
            " sold, or during periodic physical headcount audits."
        ),
        (
            "يُستخدم على الأرض فور ترقيم الماشية الجديدة، أو بيعها، أو أثناء"
            " عمليات التدقيق الدوري لتعداد الرؤوس."
        ),
    ),
    (
        "3. Birth Records / سجلات المواليد",
        (
            "A specialized log system capturing newborn lamb identification,"
            " birth weights, dam/sire linkages, and early survival tracking."
        ),
        (
            "نظام سجلات مخصص يسجل تعريف الحملان حديثة الولادة، أوزان الولادة،"
            " روابط الأم والأب، وتتبع البقاء المبكر على قيد الحياة."
        ),
        (
            "Updated daily or weekly during the lambing season right after"
            " newborn assessments in the pens."
        ),
        (
            "يتم تحديثه يومياً أو أسبوعياً خلال موسم الولادة مباشرة بعد تقييم"
            " المواليد الجدد في الحظائر."
        ),
    ),
    (
        "4. Growth Logs / سجلات النمو",
        (
            "Records sequential weight measurements to monitor development"
            " velocity, growth milestones, and feed conversion trends."
        ),
        (
            "يسجل قياسات الوزن المتسلسلة لمراقبة سرعة التطور، مراحل النمو،"
            " واتجاهات كفاءة تحويل الأعلاف."
        ),
        (
            "Utilized during scheduled monthly or quarterly weighing sessions"
            " in the farm handling yards."
        ),
        (
            "يُستخدم أثناء جلسات الوزن الشهرية أو الفصلية المجدولة في ساحات"
            " التعامل بالمزرعة."
        ),
    ),
    (
        "5. Feed Inventory / مخزون الأعلاف",
        (
            "Manages warehouse stock levels of raw and mixed feeds, tracks"
            " incoming supply deliveries, and enforces nutritional rations."
        ),
        (
            "يدير مستويات مخزون المستودعات للأعلاف الخام والمخلوطة، يتتبع عمليات"
            " تسليم الإمدادات الواردة، ويفرض الحصص الغذائية."
        ),
        (
            "Referenced daily by barn workers during feeding distribution and"
            " updated when new feed shipments arrive."
        ),
        (
            "يُرجع إليه يومياً من قبل عمال الحظائر أثناء توزيع الأعلاف ويتم"
            " تحديثه عند وصول شحنات أعلاف جديدة."
        ),
    ),
    (
        "6. Data Corrections / تصحيحات البيانات",
        (
            "An audit and error-handling utility used to correct logical"
            " discrepancies, update erroneous data inputs, or fix calculation"
            " differentials."
        ),
        (
            "أداة تدقيق ومعالجة الأخطاء تستخدم لتصحيح التناقضات المنطقية، تحديث"
            " مدخلات البيانات الخاطئة، أو إصلاح فروق الحسابات."
        ),
        (
            "Accessed immediately whenever a logical error alert or data"
            " discrepancy is flagged."
        ),
        (
            "يتم الوصول إليها فوراً كلما تم الإبلاغ عن تنبيه خطأ منطقي أو تناقض"
            " في البيانات."
        ),
    ),
    (
        "7. Performance Reports / تقارير الأداء",
        (
            "Compiles comprehensive operational analytics, flock productivity"
            " summaries, and sector-specific performance metrics."
        ),
        (
            "يجمع تحليلات تشغيلية شاملة، ملخصات إنتاجية القطيع، ومقاييس الأداء"
            " الخاصة بكل قطاع."
        ),
        (
            "Reviewed at the end of each production cycle or monthly to"
            " evaluate overall farm efficiency."
        ),
        (
            "يتم مراجعته في نهاية كل دورة إنتاجية أو شهرياً لتقييم الكفاءة العامة"
            " للمزرعة."
        ),
    ),
    (
        "8. Achievements / الإنجازات",
        (
            "Summarizes major farm milestones, production benchmarks, and"
            " target completions over time."
        ),
        (
            "يلخص المحطات الرئيسية للمزرعة، معايير الإنتاج، وإتمام الأهداف"
            " المحددة بمرور الوقت."
        ),
        (
            "Consulted periodically for stakeholder reviews or annual"
            " operational milestone evaluations."
        ),
        (
            "يُرجع إليه دورياً لمراجعات أصحاب المصلحة أو تقييمات المحطات"
            " التشغيلية السنوية."
        ),
    ),
    (
        "9. Data Audit / تدقيق البيانات",
        (
            "Evaluates database consistency, tracking integrity, and flags"
            " missing, duplicate, or irregular records across tables."
        ),
        (
            "يقيّم اتساق قاعدة البيانات، سلامة التتبع، ويحدد السجلات المفقودة أو"
            " المكررة أو غير المنتظمة عبر الجداول."
        ),
        (
            "Executed periodically by database administrators to ensure clean"
            " data hygiene."
        ),
        ("يتم تنفيذه دورياً من قبل مسؤولي قاعدة البيانات لضمان نظافة البيانات."),
    ),
    (
        "10. Breeding Prediction / تنبؤ التكاثر",
        (
            "Forecasts upcoming breeding windows and expected parturition dates"
            " based on historical flock reproductive cycles."
        ),
        (
            "يتنبأ بفترات التكاثر القادمة وتواريخ الولادة المتوقعة بناءً على دورات"
            " التكاثر التاريخية للقطيع."
        ),
        (
            "Referenced ahead of seasonal breeding schedules to prepare ram"
            " teams and maternity housing."
        ),
        ("يُرجع إليه قبل جداول التكاثر الموسمية لتجهيز فرق الكباش ومساكن" " الولادة."),
    ),
    (
        "11. Breeding Readiness / جاهزية التكاثر",
        (
            "Evaluates individual animal maturity, health metrics, and weight"
            " thresholds to screen candidates for upcoming breeding cycles."
        ),
        (
            "يقيّم نضج الحيوانات الفردية، المقاييس الصحية، وعتبات الوزن لفرش"
            " المرشحين لدورات التكاثر القادمة."
        ),
        (
            "Utilized during pre-breeding sorting sessions in the yards to"
            " select eligible ewes and rams."
        ),
        (
            "يُستخدم أثناء جلسات الفرز قبل التكاثر في الساحات لاختيار النعاج"
            " والكباش المؤهلة."
        ),
    ),
    (
        "12. Off-Take History / تاريخ الاستبعاد",
        (
            "Tracks historical sales, culling events, mortality removals, and"
            " livestock departures from the active inventory."
        ),
        (
            "يتتبع المبيعات التاريخية، عمليات الفرز/الإعدام، إزالات الوفيات،"
            " ومغادرة المواشي من المخزون النشط."
        ),
        (
            "Consulted after sales transactions or monthly mortality audits to"
            " reconcile inventory counts."
        ),
        (
            "يُرجع إليه بعد معاملات البيع أو تدقيقات الوفيات الشهرية لمطابقة"
            " أعداد المخزون."
        ),
    ),
    (
        "13. Cutting & Butcher Management / إدارة التقطيع والجزارين",
        (
            "Manages custom processing specs across three interactive tabs:"
            " registering client cutting orders (Tab 1), tracking active"
            " workflow statuses (Tab 2), and logging butcher shift productivity"
            " and task roles (Tab 3)."
        ),
        (
            "يدير مواصفات المعالجة المخصصة عبر ثلاثة تبويبات تفاعلية: تسجيل طلبات"
            " تقطيع العميل (التبويب 1)، تتبع حالات سير العمل النشطة (التبويب 2)،"
            " وتسجيل إنتاجية ورديات الجزارين وأدوار المهام (التبويب 3)."
        ),
        (
            "Used on processing days when receiving client orders, packaging"
            " meat boxes to target weights, and logging shift outputs."
        ),
        (
            "يُستخدم في أيام المعالجة عند تلقي طلبات العملاء، تعبئة صناديق"
            " اللحوم بأوزان مستهدفة، وتسجيل مخرجات ورديات الجزارين."
        ),
    ),
]

for title, func_en, func_ar, time_en, time_ar in modules_app:
    doc.add_heading(title, level=3)
    doc.add_paragraph(f"Function (English): {func_en}")
    doc.add_paragraph(f"الوظيفة (عربي): {func_ar}")
    doc.add_paragraph(f"Operational Timing (English): {time_en}")
    doc.add_paragraph(f"التوقيت التشغيلي (عربي): {time_ar}")

# Part 4: Relationship Between Applications
doc.add_heading("Part 4: Relationship Between MyHerdApp and MyHerdFinance", level=2)
doc.add_heading(
    "الجزء الرابع: العلاقة بين تطبيق MyHerdApp وتطبيق MyHerdFinance", level=2
)
doc.add_paragraph(
    "MyHerdApp acts as the operational and biological engine of the farm"
    " (recording animal counts, birth rates, growth metrics, feed consumption,"
    " and processing/cutting logistics). Meanwhile, MyHerdFinance acts as the"
    " economic mirror. Operational events in MyHerdApp—such as consuming feed"
    " stock or logging butcher labor—provide the underlying data required via"
    " the Data Bridge to calculate exact operational costs, inventory"
    " valuations, and net profitability within MyHerdFinance."
)
doc.add_paragraph(
    "عمل تطبيق MyHerdApp بمثابة المحرك التشغيلي والبيولوجي للمزرعة (تسجيل أعداد"
    " الحيوانات، معدلات الولادة، مقاييس النمو، استهلاك الأعلاف، لوجستيات"
    " المعالجة والتقطيع). وفي الوقت نفسه، يعمل تطبيق MyHerdFinance بمثابة"
    " المرآة الاقتصادية. الأحداث التشغيلية في MyHerdApp — مثل استهلاك مخزون"
    " الأعلاف أو تسجيل عمالة الجزارين — توفر البيانات الأساسية المطلوبة عبر"
    " جسر البيانات (Data Bridge) لحساب التكاليف التشغيلية بدقة، وتقييمات"
    " المخزون، وصافي الربحية داخل MyHerdFinance."
)

# Part 5: Appendix - Data Bridge Integration
doc.add_heading("Part 5: Appendix — Data Bridge Integration & Database Impact", level=2)
doc.add_heading(
    "الجزء الخامس: الملحق — تكامل جسر البيانات وتأثيره على قاعدة البيانات",
    level=2,
)

doc.add_heading(
    "1. Appendix Overview & Purpose / نظرة عامة على الملحق والغرض منه", level=3
)
doc.add_paragraph(
    "This appendix details the technical bridge connecting physical farm"
    " management operations in MyHerdApp to financial accounting in"
    " MyHerdFinance. It ensures that real-world agricultural activities"
    " translate accurately into financial ledgers."
)
doc.add_paragraph(
    "يوضح هذا الملحق الجسر التقني الذي يربط عمليات إدارة المزرعة الفيزيائية في"
    " MyHerdApp بالمحاسبة المالية في MyHerdFinance. وهو يضمن تحويل الأنشطة"
    " الزراعية في العالم الحقيقي بدقة إلى دفاتر أستاذ مالية."
)

doc.add_heading("2. Operational Workflow / سير العمل التشغيلي لجسر البيانات", level=3)
doc.add_paragraph(
    "1. Data Initiation: Physical actions occur in MyHerdApp (e.g., feed stock"
    " depletion, livestock sales, or butcher processing labor logs)."
)
doc.add_paragraph(
    "2. Selection & Filtering: The user accesses the Data Bridge module in"
    " MyHerdFinance, selecting the specific batch, date range, or activity log"
    " category."
)
doc.add_paragraph(
    "3. Preview & Verification: The system displays a translation preview"
    " showing how physical units map to monetary values."
)
doc.add_paragraph(
    "4. Execution & Posting: Upon confirmation, the bridge pushes the"
    " transaction records into the financial accounting tables."
)

doc.add_heading(
    "3. Database Impact and Table Mapping / تأثير قاعدة البيانات وربط" " الجداول",
    level=3,
)
doc.add_paragraph(
    "Source Tables (Read-Only Operational Data): cutting_orders,and"
    " butcher_performance_logs, feed_inventory & consumption logs, inventory"
    " tracking tables for livestock off-takes."
)
doc.add_paragraph(
    "Destination Tables (Write/Insert Financial Ledgers):"
    " financial_transactions, operational_outflow, revenue_inflow."
)

# Save document
doc.save("Bilingual_Administration_Manual.docx")
print("Successfully generated Bilingual_Administration_Manual.docx")

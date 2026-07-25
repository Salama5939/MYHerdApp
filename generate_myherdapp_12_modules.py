import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt as PPTxPt
from pptx.dml.color import RGBColor

# 12-Module definitions for English and Arabic
modules_en = [
    (
        "Page 1: Executive Home & Authentication Portal",
        "Operational: Secure user login, role-based access, and bilingual language toggles.\nAdministrative: System security audit logs and active session monitoring.",
    ),
    (
        "Page 2: Master Herd Roster & Inventory",
        "Operational: Individual animal tracking via unique tag numbers, breed categorization, and status filtering.\nAdministrative: Aggregate headcount reports and flock demographic summaries.",
    ),
    (
        "Page 3: Weight & Growth Tracking Staging",
        "Operational: Logging routine scale measurements and tracking individual weight gain over time.\nAdministrative: Growth curve analysis and identification of performance lags.",
    ),
    (
        "Page 4: Feed Formulation & Ration Calculator",
        "Operational: Entering daily feed formulas, calculating dry matter intake, and managing recipe mix types.\nAdministrative: Monitoring feed cost efficiency and raw material consumption rates.",
    ),
    (
        "Page 5: Health, Veterinary & Treatment Logs",
        "Operational: Recording vaccinations, medication administration, and veterinary check-ups per tag number.\nAdministrative: Disease incidence tracking and medical expense auditing.",
    ),
    (
        "Page 6: Breeding & Reproduction Register",
        "Operational: Tracking mating dates, pregnancy checks, lambing/kidding events, and litter sizes.\nAdministrative: Reproductive efficiency metrics and foundation herd productivity KPIs.",
    ),
    (
        "Page 7: Livestock Acquisition & Capital Setup",
        "Operational: Onboarding new purchase batches with source details, purchase prices, and initial weights.\nAdministrative: Capital deployment tracking and initial asset valuation.",
    ),
    (
        "Page 8: Sales & Revenue Realization",
        "Operational: Logging livestock sales, wool yields, and market transaction receipts.\nAdministrative: Revenue stream analysis and batch profitability realization.",
    ),
    (
        "Page 9: Operational Expenses & Feed Outflows (OpEx)",
        "Operational: Recording daily farm running costs, labor, and consumable purchases.\nAdministrative: Cost-center allocation and budget variance monitoring.",
    ),
    (
        "Page 10: Operational Data Bridge & Auditing",
        "Operational: Reconciling unverified field logs and processing automated feed data streams.\nAdministrative: Anomaly detection review (e.g., high-quantity consumption warnings) before ledger posting.",
    ),
    (
        "Page 11: Performance Analytics & Cost of Gain",
        "Operational: Reviewing batch-level profit and loss statements.\nAdministrative: Analyzing feed conversion ratios, cost of gain ($/kg), and operational efficiency benchmarks.",
    ),
    (
        "Page 12: Cash Flow Projections & Executive Reporting",
        "Operational: Entering scheduled future expenses and anticipated receivables.\nAdministrative: 90-day liquidity forecasting and board-level financial summaries.",
    ),
]

modules_ar = [
    (
        "الصفحة 1: البوابة الرئيسية التنفيذية وبوابة المصادقة",
        "تشغيلي: تسجيل دخول آمن للمستخدم، وصول قائم على الأدوار، وتبديل اللغات ثنائي اللغة.\nإداري: سجلات تدقيق أمان النظام ومراقبة الجلسات النشطة.",
    ),
    (
        "الصفحة 2: سجل القطيع الرئيسي والمخزون",
        "تشغيلي: تتبع الحيوانات الفردية عبر أرقام أوسم فريدة، تصنيف السلالات، وتصفية الحالة.\nإداري: تقارير تعداد الرأس الإجمالية وملخصات التركيبة السكانية للقطيع.",
    ),
    (
        "الصفحة 3: تتبع الأوزان ونمو الحيوانات",
        "تشغيلي: تسجيل قياسات الموازين الروتينية وتتبع زيادة الوزن الفردية بمرور الوقت.\nإداري: تحليل منحنيات النمو وتحديد التأخيرات في الأداء.",
    ),
    (
        "الصفحة 4: تركيبة الأعلاف وحاسبة الحصص",
        "تشغيلي: إدخال وصفات الأعلاف اليومية، حساب استهلاك المادة الجافة، وإدارة أنواع خلطات الوصفات.\nإداري: مراقبة كفاءة تكلفة الأعلاف ومعدلات استهلاك المواد الخام.",
    ),
    (
        "الصفحة 5: سجلات الصحة والبيطرة والعلاج",
        "تشغيلي: تسجيل التطعيمات، إعطاء الأدوية، والفحوصات البيطرية لكل رقم وسم.\nإداري: تتبع معدلات انتشار الأمراض وتدقيق النفقات الطبية.",
    ),
    (
        "الصفحة 6: سجل التكاثر والتناسل",
        "تشغيلي: تتبع تواريخ التزاوج، فحوصات الحمل، أحداث الولادة، وأحجام البطون.\nإداري: مقاييس الكفاءة التناسلية ومؤشرات الأداء الرئيسية لإنتاجية القطيع الأساسي.",
    ),
    (
        "الصفحة 7: حيازة المواشي وإعداد رأس المال",
        "تشغيلي: إدراج دفعات الشراء الجديدة مع تفاصيل المصدر، أسعار الشراء، والأوزان الأولية.\nإداري: تتبع نشر رأس المال وتقييم الأصول الأولية.",
    ),
    (
        "الصفحة 8: المبيعات وتحقيق الإيرادات",
        "تشغيلي: تسجيل مبيعات المواشي، عوائد الصوف، وإيصالات معاملات السوق.\nإداري: تحليل تدفقات الإيرادات وتحقيق ربحية الدفعات.",
    ),
    (
        "الصفحة 9: المصروفات التشغيلية وتدفقات الأعلاف (OpEx)",
        "تشغيلي: تسجيل تكاليف تشغيل المزرعة اليومية، العمالة، ومشتريات المواد المستهلكة.\nإداري: تخصيص مراكز التكلفة ومراقبة انحراف الميزانية.",
    ),
    (
        "الصفحة 10: جسر البيانات التشغيلية والتدقيق",
        "تشغيلي: مطابقة السجلات الميدانية غير المتحقق منها ومعالجة تدفقات بيانات الأعلاف التلقائية.\nإداري: مراجعة كشف الشذوذ (مثل تحذيرات الاستهلاك العالي) قبل الترحيل للدفتر.",
    ),
    (
        "الصفحة 11: تحليلات الأداء وتكلفة زيادة الوزن",
        "تشغيلي: مراجعة قوائم الأرباح والخسائر على مستوى الدفعات.\nإداري: تحليل نسب تحويل الأعلاف، تكلفة زيادة الوزن ($/كجم)، ومعايير الكفاءة التشغيلية.",
    ),
    (
        "الصفحة 12: توقعات التدفق النقدي والتقارير التنفيذية",
        "تشغيلي: إدخال المصروفات المستقبلية المجدولة والمتحصلات المتوقعة.\nإداري: التنبؤ بالسيولة لـ 90 يوماً والملخصات المالية على مستوى مجلس الإدارة.",
    ),
]


def generate_word_documents():
    # English Word Manual
    doc_en = docx.Document()
    t_en = doc_en.add_heading(
        "MyHerdApp - Comprehensive 12-Module Reference Manual", level=0
    )
    t_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_en = doc_en.add_paragraph(
        "Complete Operational Workflows & Administrative Oversight Guide"
    )
    s_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_en.runs[0].font.italic = True

    doc_en.add_heading("System Architecture & 12-Page Reference", level=1)
    for title, desc in modules_en:
        p = doc_en.add_paragraph(style="List Bullet")
        p.add_run(f"{title}\n").bold = True
        p.add_run(desc)

    doc_en.save("MyHerdApp_12_Module_Manual_English.docx")

    # Arabic Word Manual
    doc_ar = docx.Document()
    t_ar = doc_ar.add_heading(
        "تطبيق إدارة القطيع (MyHerdApp) - الدليل المرجعي الشامل لـ 12 وحدة", level=0
    )
    t_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_ar = doc_ar.add_paragraph("دليل سير العمل التشغيلي والإشراف الإداري المتكامل")
    s_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_ar.runs[0].font.italic = True

    doc_ar.add_heading("هندسة النظام والمرجع لـ 12 صفحة", level=1)
    for title, desc in modules_ar:
        p = doc_ar.add_paragraph(style="List Bullet")
        p.add_run(f"{title}\n").bold = True
        p.add_run(desc)

    doc_ar.save("MyHerdApp_12_Module_Manual_Arabic.docx")
    print("Word manuals generated successfully.")


def generate_powerpoint_presentations():
    def set_slide_title(slide, text):
        title_shape = slide.shapes.title
        if title_shape is not None:
            title_shape.text = text
            return title_shape
        if len(slide.placeholders) > 0:
            slide.placeholders[0].text = text
            return slide.placeholders[0]
        return None

    # English Presentation
    prs_en = Presentation()
    prs_en.slide_width = Inches(13.333)
    prs_en.slide_height = Inches(7.5)

    # Title slide
    s1_en = prs_en.slides.add_slide(prs_en.slide_layouts[0])
    title_shape_en = set_slide_title(
        s1_en, "MyHerdApp 12-Module Executive Presentation"
    )
    # Safely set placeholder text using text_frame (avoids direct BaseShape.text assignment)
    placeholder_en = s1_en.placeholders[1]
    text_frame_en = getattr(placeholder_en, "text_frame", None)
    if text_frame_en is not None:
        text_frame_en.clear()
        if text_frame_en.paragraphs:
            text_frame_en.paragraphs[0].text = (
                "Comprehensive Livestock Operations & Administrative Oversight"
            )
    else:
        try:
            # fallback: use text_frame if available at runtime
            tf_fallback = getattr(placeholder_en, "text_frame", None)
            if tf_fallback is not None:
                tf_fallback.clear()
                if tf_fallback.paragraphs:
                    tf_fallback.paragraphs[0].text = (
                        "Comprehensive Livestock Operations & Administrative Oversight"
                    )
        except Exception:
            pass

    for title, desc in modules_en:
        slide = prs_en.slides.add_slide(prs_en.slide_layouts[1])
        set_slide_title(slide, title)
        placeholder = slide.placeholders[1]
        tf = getattr(placeholder, "text_frame", None)
        if tf is not None:
            tf.clear()
            for idx, line in enumerate(desc.split("\n")):
                p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
                p.text = f"• {line}"
                p.font.size = Pt(20)
                p.space_after = Pt(14)

    prs_en.save("MyHerdApp_12_Module_Presentation_English.pptx")

    # Arabic Presentation
    prs_ar = Presentation()
    prs_ar.slide_width = Inches(13.333)
    prs_ar.slide_height = Inches(7.5)

    # Title slide
    s1_ar = prs_ar.slides.add_slide(prs_ar.slide_layouts[0])
    title_shape_ar = set_slide_title(
        s1_ar, "عرض مجلس الإدارة لـ 12 وحدة لتطبيق MyHerdApp"
    )
    if title_shape_ar is not None and hasattr(title_shape_ar, "text_frame"):
        title_shape_ar_text_frame = getattr(title_shape_ar, "text_frame")
        for p in title_shape_ar_text_frame.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder_ar = s1_ar.placeholders[1]
    if hasattr(placeholder_ar, "text_frame"):
        placeholder_ar_text_frame = getattr(placeholder_ar, "text_frame")
        placeholder_ar_text_frame.text = (
            "عمليات الثروة الحيوانية الشاملة والإشراف الإداري"
        )
        for p in placeholder_ar_text_frame.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for title, desc in modules_ar:
        slide = prs_ar.slides.add_slide(prs_ar.slide_layouts[1])
        title_shape_ar = set_slide_title(slide, title)
        if title_shape_ar is not None and hasattr(title_shape_ar, "text_frame"):
            title_shape_ar_text_frame = getattr(title_shape_ar, "text_frame")
            for p in title_shape_ar_text_frame.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        placeholder_ar_body = slide.placeholders[1]
        tf = getattr(placeholder_ar_body, "text_frame", None)
        if tf is not None:
            tf.clear()
            for idx, line in enumerate(desc.split("\n")):
                p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
                p.text = f"• {line}"
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.font.size = Pt(20)
                p.space_after = Pt(14)

    prs_ar.save("MyHerdApp_12_Module_Presentation_Arabic.pptx")
    print("PowerPoint decks generated successfully.")


if __name__ == "__main__":
    generate_word_documents()
    generate_powerpoint_presentations()
    print("All 12-module documents and presentations created successfully!")

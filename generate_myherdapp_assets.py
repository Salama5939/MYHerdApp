import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt as PPTxPt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def generate_word_docs():
    # --- 1. ENGLISH WORD MANUAL ---
    doc_en = docx.Document()
    title_en = doc_en.add_heading("MyHerdApp - Dual-Purpose Reference Manual", level=0)
    title_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_en = doc_en.add_paragraph("Operational Workflows & Livestock Management Guide")
    sub_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_en.runs[0].font.italic = True

    doc_en.add_heading("Part 1: Farm Staff Operational Manual", level=1)
    doc_en.add_paragraph(
        "Objective: Step-by-step guidelines for daily livestock tracking, weight logging, and feed formula execution."
    )

    staff_bullets_en = [
        (
            "Tag Registration:",
            "Register new sheep into the system using unique tag numbers (e.g., Tag #114).",
        ),
        (
            "Weight Logging:",
            "Record routine weight logs to monitor individual animal growth and fattening progress.",
        ),
        (
            "Feed Formulation:",
            "Apply daily feed formulas and manage recipe mixes for specific herd groups.",
        ),
    ]
    for heading, text in staff_bullets_en:
        p = doc_en.add_paragraph(style="List Bullet")
        p.add_run(heading).bold = True
        p.add_run(f" {text}")

    doc_en.add_heading("Part 2: Administrative & Management Manual", level=1)
    admin_bullets_en = [
        (
            "Herd Logistics:",
            "Monitor overall flock distribution, active statuses (Fattening vs. Foundation), and inventory metrics.",
        ),
        (
            "Performance Oversight:",
            "Analyze weight gains, feed conversion rates, and historical logs across batches.",
        ),
        (
            "System Deployment:",
            "Manage cloud-based access (Flask/Render architecture) and database stability.",
        ),
    ]
    for heading, text in admin_bullets_en:
        p = doc_en.add_paragraph(style="List Bullet")
        p.add_run(heading).bold = True
        p.add_run(f" {text}")

    doc_en.save("MyHerdApp_Reference_Manual_English.docx")

    # --- 2. ARABIC WORD MANUAL ---
    doc_ar = docx.Document()
    title_ar = doc_ar.add_heading(
        "دليل المرجعية لتطبيق إدارة القطيع (MyHerdApp)", level=0
    )
    title_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_ar = doc_ar.add_paragraph("سير العمل التشغيلي ودليل إدارة الثروة الحيوانية")
    sub_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_ar.runs[0].font.italic = True

    doc_ar.add_heading("الجزء الأول: دليل تشغيل طاقم المزرعة", level=1)
    doc_ar.add_paragraph(
        "الهدف: إرشادات خطوة بخطوة لتتبع الماشية اليومي، تسجيل الأوزان، وتنفيذ وصفات الأعلاف."
    )

    staff_bullets_ar = [
        (
            "تسجيل الأوسم:",
            "تسجيل الأغنام الجديدة في النظام باستخدام أرقام أوسم فريدة (مثال: الوسم رقم 114).",
        ),
        (
            "تسجيل الأوزان:",
            "تسجيل أوزان روتينية لمراقبة نمو الحيوانات الفردية وتقدم التسمين.",
        ),
        (
            "تركيب الأعلاف:",
            "تطبيق خلطات الأعلاف اليومية وإدارة وصفات الخلط لمجموعات القطيع المحددة.",
        ),
    ]
    for heading, text in staff_bullets_ar:
        p = doc_ar.add_paragraph(style="List Bullet")
        p.add_run(heading).bold = True
        p.add_run(f" {text}")

    doc_ar.add_heading("الجزء الثاني: الدليل الإداري وإدارة القطيع", level=1)
    admin_bullets_ar = [
        (
            "لوجستيات القطيع:",
            "مراقبة توزع القطيع الإجمالي، الحالات النشطة (التسمين مقابل الأساسي)، ومقاييس المخزون.",
        ),
        (
            "رقابة الأداء:",
            "تحليل زيادة الوزن، معدلات تحويل الأعلاف، والسجلات التاريخية عبر الدفعات.",
        ),
        (
            "نشر النظام:",
            "إدارة الوصول السحابي (هندسة Flask/Render) واستقرار قاعدة البيانات.",
        ),
    ]
    for heading, text in admin_bullets_ar:
        p = doc_ar.add_paragraph(style="List Bullet")
        p.add_run(heading).bold = True
        p.add_run(f" {text}")

    doc_ar.save("MyHerdApp_Reference_Manual_Arabic.docx")
    print("Word documents generated successfully!")


def generate_powerpoint_decks():
    # --- 3. ENGLISH PRESENTATION ---
    prs_en = Presentation()
    prs_en.slide_width = Inches(13.333)
    prs_en.slide_height = Inches(7.5)

    def add_ppt_slide(prs, title_text, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        title.text = title_text
        for p in title.text_frame.paragraphs:
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(24, 43, 73)
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, item in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = item
            p.font.size = Pt(20)
            p.space_after = Pt(14)

    # Title Slide Eng
    slide1 = prs_en.slides.add_slide(prs_en.slide_layouts[0])
    # Safely set title and subtitle (some slide layouts may not have these placeholders)
    if slide1.shapes.title is not None:
        slide1.shapes.title.text = "MyHerdApp Board Review"
    else:
        txBox = slide1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
        tf = txBox.text_frame
        tf.text = "MyHerdApp Board Review"
        for p in tf.paragraphs:
            p.font.size = PPTxPt(32)
    try:
        ph = slide1.placeholders[1]
        # some placeholder shapes may not expose text_frame at type-check time;
        # guard with hasattr and has_text_frame before accessing
        if ph is not None and getattr(ph, "has_text_frame", False):
            text_frame = getattr(ph, "text_frame", None)
            if text_frame is not None:
                text_frame.paragraphs[0].text = (
                    "Livestock Management & Operational Oversight"
                )
    except Exception:
        # fallback: add a textbox for subtitle
        txBox2 = slide1.shapes.add_textbox(
            Inches(1), Inches(1.4), Inches(11), Inches(1)
        )
        tf2 = txBox2.text_frame
        tf2.text = "Livestock Management & Operational Oversight"

    add_ppt_slide(
        prs_en,
        "1. Executive Herd Scale & Overview",
        [
            "• Core Objective: Streamlining sheep farming operations and tracking individual livestock metrics.",
            "• Flock Visibility: Real-time tracking of tag numbers, weight progression, and physiological status.",
            "• Operational Efficiency: Digitalizing daily feed formulas and animal logistics.",
        ],
    )
    add_ppt_slide(
        prs_en,
        "2. Livestock Performance Analytics",
        [
            "• Weight Gain Tracking: Monitoring maximum and minimum weight thresholds across batches.",
            "• Feed Optimization: Evaluating daily ration metrics and recipe costs.",
            "• Growth Monitoring: Identifying high-performing fattening cycles and addressing growth lags.",
        ],
    )
    add_ppt_slide(
        prs_en,
        "3. Cloud Architecture & Deployment",
        [
            "• Tech Stack: Built with Python Flask, SQLite, and deployed on Render cloud services.",
            "• Accessibility: Seamless verification and operation across desktop and mobile devices.",
            "• Data Integrity: Secure repository maintenance and automated backup structures.",
        ],
    )
    prs_en.save("MyHerdApp_Board_Presentation_English.pptx")

    # --- 4. ARABIC PRESENTATION ---
    prs_ar = Presentation()
    prs_ar.slide_width = Inches(13.333)
    prs_ar.slide_height = Inches(7.5)

    def add_ppt_slide_ar(prs, title_text, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        title.text = title_text
        for p in title.text_frame.paragraphs:
            p.alignment = PP_ALIGN.RIGHT
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(24, 43, 73)
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, item in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = item
            p.alignment = PP_ALIGN.RIGHT
            p.font.size = Pt(20)
            p.space_after = Pt(14)

    # Title Slide Ar
    slide1_ar = prs_ar.slides.add_slide(prs_ar.slide_layouts[0])
    if slide1_ar.shapes.title is not None:
        slide1_ar.shapes.title.text = "عرض مجلس الإدارة لتطبيق MyHerdApp"
        for p in slide1_ar.shapes.title.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    else:
        txBox_ar = slide1_ar.shapes.add_textbox(
            Inches(1), Inches(0.5), Inches(11), Inches(1)
        )
        tf_ar = txBox_ar.text_frame
        tf_ar.text = "عرض مجلس الإدارة لتطبيق MyHerdApp"
        for p in tf_ar.paragraphs:
            p.alignment = PP_ALIGN.CENTER
    try:
        ph_ar = slide1_ar.placeholders[1]
        # guard against placeholder shapes that may not expose text_frame
        if ph_ar is not None and getattr(ph_ar, "has_text_frame", False):
            text_frame_ar = getattr(ph_ar, "text_frame", None)
            if text_frame_ar is not None:
                text_frame_ar.paragraphs[0].text = (
                    "إدارة الثروة الحيوانية والإشراف التشغيلي"
                )
                for p in text_frame_ar.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
    except Exception:
        txBox2_ar = slide1_ar.shapes.add_textbox(
            Inches(1), Inches(1.4), Inches(11), Inches(1)
        )
        tf2_ar = txBox2_ar.text_frame
        tf2_ar.text = "إدارة الثروة الحيوانية والإشراف التشغيلي"
        for p in tf2_ar.paragraphs:
            p.alignment = PP_ALIGN.CENTER

    add_ppt_slide_ar(
        prs_ar,
        "1. نطاق القطيع ونظرة عامة",
        [
            "• الهدف الرئيسي: تبسيط عمليات تربية الأغنام وتتبع مقاييس الثروة الحيوانية الفردية.",
            "• رؤية القطيع: تتبع فوري لأرقام الأوسم، تقدم الأوزان، والحالة الفسيولوجية.",
            "• الكفاءة التشغيلية: رقمنة وصفات الأعلاف اليومية ولوجستيات الحيوانات.",
        ],
    )
    add_ppt_slide_ar(
        prs_ar,
        "2. تحليلات أداء الماشية",
        [
            "• تتبع زيادة الوزن: مراقبة الحد الأقصى والأدنى لأوزان الحيوانات عبر الدفعات.",
            "• تحسين الأعلاف: تقييم مقاييس الحصص اليومية وتكاليف الوصفات.",
            "• مراقبة النمو: تحديد دورات التسمين عالية الأداء ومعالجة التأخيرات.",
        ],
    )
    add_ppt_slide_ar(
        prs_ar,
        "3. البنية السحابية والنشر",
        [
            "• التقنيات المستخدمة: مبني بلغة بايثون Flask وقاعدة بيانات SQLite ومنشور سحابياً عبر Render.",
            "• سهولة الوصول: التحقق السلس والتشغيل عبر أجهزة سطح المكتب والأجهزة المحمولة.",
            "• سلامة البيانات: صيانة آمنة لمستودع الكود البرمجي وهياكل النسخ الاحتياطي التلقائي.",
        ],
    )
    prs_ar.save("MyHerdApp_Board_Presentation_Arabic.pptx")
    print("PowerPoint presentations generated successfully!")


if __name__ == "__main__":
    generate_word_docs()
    generate_powerpoint_decks()
    print(
        "All MyHerdApp documents and presentations (English & Arabic) created successfully!"
    )
